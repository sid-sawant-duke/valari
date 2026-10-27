"""Builds the Google Sheets connector template (.xlsx, import into Google Sheets).

Usage: python3 scripts/build-sheet.py sheet/valari-template.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]
FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F2937")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", start_color="FFF7D6")

wb = Workbook()

readme = wb.active
readme.title = "Read Me"
lines = [
    ("Valari: Google Sheets connector template", True),
    ("", False),
    ("Use this if you don't use Salesforce, or as a product-user list next to your CRM.", False),
    ("", False),
    ("How to use", True),
    ("1. Import into Google Sheets: File > Import > Upload > Replace spreadsheet.", False),
    ("2. Champions: one row per past user of your product. Yellow cells are yours to fill in.", False),
    ("3. Customers: companies you already have a relationship with, keyed by email domain. If a champion moves to one of these, that account's owner is alerted and the champion's own AE is cc'd.", False),
    ("4. Moves: leave empty. The connector adds a row for every confirmed move.", False),
    ("5. The example rows have do_not_contact = TRUE, so they are skipped. Delete them when your real rows are in.", False),
    ("6. Paste this sheet's ID into the connector's Config node (from the URL: /spreadsheets/d/ID/edit).", False),
    ("", False),
    ("Champions columns", True),
    ("contact_id: any unique ID. Optional for product-user lists (the email is used instead).", False),
    ("work_email: their email at the old company. Bounces and auto-replies are matched on this. Required.", False),
    ("linkedin_url: the most reliable way to look someone up. Strongly recommended.", False),
    ("role_type: buyer, champion or user. Scoring gives buyers the most points.", False),
    ("owner_ae / owner_ae_email / owner_ae_slack_id: who owns this person. Or leave blank and use the team directory in the core Config.", False),
    ("do_not_contact: TRUE skips the person entirely (unhappy churn, opted out, etc.).", False),
    ("", False),
    ("Tracker state (status, rechecks, cases) is kept in n8n Data Tables, not in this sheet.", True),
    ("Rename columns only if you also update the Map Champions / Map Accounts nodes in the connector.", True),
]
for i, (text, bold) in enumerate(lines, start=1):
    cell = readme.cell(row=i, column=1, value=text)
    cell.font = Font(name=FONT, bold=bold, size=14 if i == 1 else 11)
readme.column_dimensions["A"].width = 125


def tab(title, columns, rows, user_cols, widths, validations=None, notes=None):
    ws = wb.create_sheet(title)
    for j, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=j, value=col)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[cell.column_letter].width = widths.get(col, 18)
        if notes and col in notes:
            cell.comment = Comment(notes[col], "Valari")
    for r, row in enumerate(rows, start=2):
        for j, col in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=j, value=row.get(col, ""))
            cell.font = Font(name=FONT)
            if col in user_cols:
                cell.fill = INPUT_FILL
    ws.freeze_panes = "B2"
    for col, options in (validations or {}).items():
        letter = ws.cell(row=1, column=columns.index(col) + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.add(f"{letter}2:{letter}5000")
        ws.add_data_validation(dv)


champ_cols = ["contact_id", "full_name", "work_email", "linkedin_url", "last_company", "last_company_domain", "last_title",
              "role_type", "owner_ae", "owner_ae_email", "owner_ae_slack_id", "do_not_contact", "notes"]
examples = [
    ("EX-1", "Jordan Avery", "jordan.avery@northwind.example", "https://www.linkedin.com/in/example-jordan-avery", "Northwind Logistics",
     "northwind.example", "Senior UX Researcher", "champion", "Priya Shah", "priya@yourcompany.example", "U0000000001",
     "Example row. Ran the research repository rollout to 3 teams."),
    ("EX-2", "Sam Okafor", "sam.okafor@brightcart.example", "", "Brightcart", "brightcart.example", "Director of Design",
     "buyer", "Marcus Lee", "marcus@yourcompany.example", "U0000000002", "Example row. Signed the 2025 renewal."),
]
champ_rows = [dict(zip(champ_cols[:11] + ["notes"], e), do_not_contact="TRUE") for e in examples]
tab("Champions", champ_cols, champ_rows, set(champ_cols),
    widths={"full_name": 20, "work_email": 36, "linkedin_url": 44, "last_company": 22, "last_company_domain": 26,
            "last_title": 24, "owner_ae_email": 28, "owner_ae_slack_id": 18, "notes": 50},
    validations={"role_type": ["buyer", "champion", "user"], "do_not_contact": ["TRUE", "FALSE"]},
    notes={"owner_ae_slack_id": "Slack profile > ... > Copy member ID. Used for @mentions."})

cust_cols = ["domain", "company_name", "relationship", "owner_ae", "owner_ae_email", "owner_ae_slack_id"]
cust_rows = [
    {"domain": "acme.example", "company_name": "Acme (example)", "relationship": "customer", "owner_ae": "Priya Shah", "owner_ae_email": "priya@yourcompany.example", "owner_ae_slack_id": "U0000000001"},
    {"domain": "lumen.example", "company_name": "Lumen Retail (example)", "relationship": "open_deal", "owner_ae": "Marcus Lee", "owner_ae_email": "marcus@yourcompany.example", "owner_ae_slack_id": "U0000000002"},
    {"domain": "fieldstone.example", "company_name": "Fieldstone (example)", "relationship": "lost_deal", "owner_ae": "Marcus Lee", "owner_ae_email": "marcus@yourcompany.example", "owner_ae_slack_id": "U0000000002"},
]
tab("Customers", cust_cols, cust_rows, set(cust_cols), widths={"domain": 26, "company_name": 26, "owner_ae_email": 28, "owner_ae_slack_id": 18},
    validations={"relationship": ["customer", "open_deal", "lost_deal"]}, notes={"domain": "Website domain without www, e.g. acme.com"})

move_cols = ["recorded_at", "contact_id", "full_name", "old_company", "new_company", "new_company_domain", "new_title", "new_email", "summary"]
tab("Moves", move_cols, [], set(), widths={"recorded_at": 24, "full_name": 20, "old_company": 22, "new_company": 22, "new_company_domain": 24, "new_title": 24, "new_email": 30, "summary": 80})

wb.save(OUT)
print("saved", OUT)
